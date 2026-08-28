"""PDF and DOCX normalization with stable document locators."""

from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO

from docx import Document
from pypdf import PdfReader

from ..agent_models import ContentSegment, NormalizedSource, PageRangeLocator, ParagraphRangeLocator, SectionLocator


class DocumentExtractionError(ValueError):
    pass


def _segment(identifier: int, text: str, locator: object) -> ContentSegment:
    value = text.strip()
    return ContentSegment(id=f"seg-{identifier}", text=value, locator=locator, digest=sha256(value.encode()).hexdigest())


def extract_pdf(source_id: str, title: str, body: bytes, *, receipt_id: str = "local-extraction") -> NormalizedSource:
    reader = PdfReader(BytesIO(body))
    if reader.is_encrypted:
        raise DocumentExtractionError("encrypted PDFs are unsupported")
    segments: list[ContentSegment] = []
    for page_number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip() or f"[Blank page {page_number}]"
        segments.append(_segment(len(segments) + 1, text, PageRangeLocator(startPage=page_number, endPage=page_number)))
    if not segments:
        raise DocumentExtractionError("PDF contains no pages")
    return NormalizedSource(sourceId=source_id, sourceKind="document", title=title, mimeType="application/pdf", contentDigest=sha256(body).hexdigest(), extractorVersion="pdf-pypdf-6", extractedAt=datetime.now(timezone.utc), segments=segments, metadata={"pageCount": len(reader.pages)}, extractionReceiptId=receipt_id)


def extract_docx(source_id: str, title: str, body: bytes, *, receipt_id: str = "local-extraction") -> NormalizedSource:
    document = Document(BytesIO(body))
    segments: list[ContentSegment] = []
    headings: dict[str, int] = {}
    paragraph_number = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        paragraph_number += 1
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            headings[text] = headings.get(text, 0) + 1
            locator = SectionLocator(heading=text, occurrence=headings[text])
        else:
            locator = ParagraphRangeLocator(startParagraph=paragraph_number, endParagraph=paragraph_number)
        segments.append(_segment(len(segments) + 1, text, locator))
    for table in document.tables:
        for row in table.rows:
            paragraph_number += 1
            text = " | ".join(cell.text.strip() for cell in row.cells)
            if text.strip(" |"):
                segments.append(_segment(len(segments) + 1, text, ParagraphRangeLocator(startParagraph=paragraph_number, endParagraph=paragraph_number)))
    if not segments:
        raise DocumentExtractionError("DOCX contains no extractable content")
    return NormalizedSource(sourceId=source_id, sourceKind="document", title=title, mimeType="application/vnd.openxmlformats-officedocument.wordprocessingml.document", contentDigest=sha256(body).hexdigest(), extractorVersion="docx-python-docx-1", extractedAt=datetime.now(timezone.utc), segments=segments, metadata={"paragraphCount": paragraph_number}, extractionReceiptId=receipt_id)
