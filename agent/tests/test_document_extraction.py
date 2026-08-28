from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from harmonia_agent.extraction.documents import DocumentExtractionError, extract_docx, extract_pdf


def test_pdf_preserves_one_based_pages() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    data = BytesIO()
    writer.write(data)
    result = extract_pdf("pdf-1", "Brief", data.getvalue())
    assert result.segments[0].locator.kind == "page_range"
    assert result.segments[0].locator.startPage == 1
    assert result.extractorVersion.startswith("pdf-")


def test_encrypted_pdf_is_rejected() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.encrypt("secret")
    data = BytesIO()
    writer.write(data)
    with pytest.raises(DocumentExtractionError, match="encrypted"):
        extract_pdf("pdf-1", "Protected", data.getvalue())


def test_docx_preserves_headings_paragraphs_and_tables() -> None:
    document = Document()
    document.add_heading("Launch", level=1)
    document.add_paragraph("Proof point")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "42%"
    data = BytesIO()
    document.save(data)
    result = extract_docx("docx-1", "Brief", data.getvalue())
    assert any(segment.locator.kind == "section" for segment in result.segments)
    assert "42%" in " ".join(segment.text for segment in result.segments)
